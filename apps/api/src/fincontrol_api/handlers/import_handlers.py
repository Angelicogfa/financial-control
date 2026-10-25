"""B3 statement import: preview, then commit the new lines atomically (item 12.3)."""

from __future__ import annotations

import csv
import io
from collections.abc import Callable
from datetime import date, timedelta
from decimal import Decimal

from fincontrol_core.catalog.etf_exposure import suggest_arca_class
from fincontrol_core.domain.entities import (
    AssetEntity,
    CorporateEventEntity,
    EarningEntity,
    TransactionEntity,
)
from fincontrol_core.errors import ValidationError
from fincontrol_core.imports.b3 import (
    ImportLine,
    ImportPreview,
    LineAction,
    LineStatus,
    classify,
    event_factor,
    guess_asset_type,
    parse_statement,
)
from fincontrol_core.portfolio.ledger import LedgerConflictError, validate_ledger
from fincontrol_core.portfolio.positions import InsufficientQuantityError, build_position
from fincontrol_core.ports.unit_of_work import AbstractUnitOfWork

from fincontrol_api.handlers.account_handlers import resolve_account_id
from fincontrol_api.schemas import (
    ImportCommitDTO,
    ImportLineDTO,
    ImportPreviewDTO,
    UnknownTickerDTO,
)

UoWFactory = Callable[[], AbstractUnitOfWork]
MAX_STATEMENT_BYTES = 5 * 1024 * 1024


def read_statement_rows(filename: str, content: bytes) -> list[dict[str, object]]:
    """Rows of the first sheet (.xlsx) or of a CSV export, keyed by the header row."""
    if len(content) > MAX_STATEMENT_BYTES:
        raise ValidationError("Statement larger than 5 MB", code="FILE_TOO_LARGE")
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        except Exception as exc:
            raise ValidationError("Unreadable spreadsheet", code="INVALID_FILE") from exc
        values = [list(r) for r in sheet.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else ",", delimiters=";,\t")
        values = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ValidationError("Send the .xlsx (or .csv) exported by B3", code="INVALID_FILE")
    if not values:
        raise ValidationError("Empty statement", code="INVALID_FILE")
    headers = [str(h or "").strip() for h in values[0]]
    return [dict(zip(headers, row, strict=False)) for row in values[1:]]


def _parse(filename: str, content: bytes) -> ImportPreview:
    rows = read_statement_rows(filename, content)
    try:
        return parse_statement(rows)
    except ValueError as exc:
        raise ValidationError(str(exc), code="INVALID_STATEMENT") from exc


def _line_dto(line: ImportLine) -> ImportLineDTO:
    return ImportLineDTO(
        line_no=line.line_no,
        action=line.action.value,
        status=line.status.value,
        entry_date=line.when,
        ticker=line.ticker,
        description=line.description,
        quantity=line.quantity,
        unit_price=line.unit_price,
        amount=line.amount,
        trade_kind=line.trade_kind.value if line.trade_kind else None,
        earning_kind=line.earning_kind.value if line.earning_kind else None,
        event_kind=line.event_kind.value if line.event_kind else None,
        note=line.note,
    )


def _held_before(assets: dict[str, AssetEntity]) -> Callable[[str, date], Decimal]:
    def held(ticker: str, when: date) -> Decimal:
        asset = assets.get(ticker)
        if asset is None:
            return Decimal("0")
        try:
            return build_position(
                asset.transactions,
                up_to_date=when - timedelta(days=1),
                events=asset.corporate_events,
            ).quantity
        except InsufficientQuantityError:
            return Decimal("0")

    return held


class ImportHandler:
    """Preview shows what would happen; commit applies only NEW lines, all or nothing."""

    def __init__(self, uow_factory: UoWFactory) -> None:
        self.uow_factory = uow_factory

    def preview(self, filename: str, content: bytes) -> ImportPreviewDTO:
        parsed = _parse(filename, content)
        with self.uow_factory() as uow:
            assets = {a.ticker: a for a in uow.assets.list_all(active_only=False)}
            preview = classify(
                parsed,
                known_tickers=set(assets),
                existing_hashes=uow.ledger_imports.existing_hashes(
                    [line.hash for line in parsed.lines]
                ),
                held_before=_held_before(assets),
            )
            unknown = []
            for ticker in preview.unknown_tickers:
                listed = uow.catalog.get(ticker)
                unknown.append(
                    UnknownTickerDTO(
                        ticker=ticker,
                        name=listed.name if listed else None,
                        asset_type=(listed.asset_type if listed else None)
                        or guess_asset_type(ticker),
                    )
                )
        return ImportPreviewDTO(
            kind=preview.kind.value,
            counts=preview.counts,
            unknown_tickers=unknown,
            lines=[_line_dto(line) for line in preview.lines],
        )

    def commit(
        self,
        filename: str,
        content: bytes,
        *,
        account_id: int | None,
        create_tickers: list[str],
    ) -> ImportCommitDTO:
        parsed = _parse(filename, content)
        with self.uow_factory() as uow:
            account = resolve_account_id(uow, account_id)
            for ticker in sorted({t.strip().upper() for t in create_tickers if t.strip()}):
                if uow.assets.get_by_ticker(ticker) is not None:
                    continue
                listed = uow.catalog.get(ticker)
                asset_type = (listed.asset_type if listed else None) or guess_asset_type(ticker)
                if asset_type is None:
                    raise ValidationError(
                        f"Cannot tell the type of {ticker}; register it first",
                        code="UNKNOWN_INSTRUMENT",
                    )
                uow.assets.add(
                    AssetEntity(
                        ticker=ticker,
                        name=listed.name if listed else ticker,
                        asset_type=asset_type,
                        arca_class=suggest_arca_class(ticker, asset_type).arca_class,
                    )
                )

            assets = {a.ticker: a for a in uow.assets.list_all(active_only=False)}
            preview = classify(
                parsed,
                known_tickers=set(assets),
                existing_hashes=uow.ledger_imports.existing_hashes(
                    [line.hash for line in parsed.lines]
                ),
            )
            created = self._apply(uow, preview.new_lines, assets, account)
            uow.commit()
        return ImportCommitDTO(
            kind=preview.kind.value,
            created=created,
            skipped={k: v for k, v in preview.counts.items() if k != LineStatus.NEW.value},
            created_assets=sorted(
                t.strip().upper() for t in create_tickers if t.strip().upper() in assets
            ),
        )

    @staticmethod
    def _apply(
        uow: AbstractUnitOfWork,
        lines: list[ImportLine],
        assets: dict[str, AssetEntity],
        account_id: int | None,
    ) -> dict[str, int]:
        created = {"TRADE": 0, "EARNING": 0, "EVENT": 0}
        by_ticker: dict[str, list[ImportLine]] = {}
        for line in lines:
            by_ticker.setdefault(line.ticker or "", []).append(line)

        for ticker, items in sorted(by_ticker.items()):
            asset = assets[ticker]
            asset_id = asset.id or 0
            note = "Importado da B3 (linha {})"
            trades = [
                TransactionEntity(
                    asset_id=asset_id,
                    kind=line.trade_kind,  # type: ignore[arg-type]
                    trade_date=line.when,  # type: ignore[arg-type]
                    quantity=line.quantity or Decimal("0"),
                    unit_price=line.unit_price or Decimal("0"),
                    notes=note.format(line.line_no),
                    account_id=account_id,
                    import_hash=line.hash,
                )
                for line in items
                if line.action == LineAction.TRADE
            ]
            ledger = [*asset.transactions, *trades]
            events = list(asset.corporate_events)
            new_events = []
            for line in sorted(
                (i for i in items if i.action == LineAction.EVENT),
                key=lambda i: (i.when, i.line_no),
            ):
                when = line.when or date.min
                try:
                    held = build_position(
                        ledger, up_to_date=when - timedelta(days=1), events=events
                    ).quantity
                except InsufficientQuantityError:
                    held = Decimal("0")
                factor = event_factor(line.event_kind, line.quantity or Decimal("0"), held)  # type: ignore[arg-type]
                if factor is None:
                    raise ValidationError(
                        f"{ticker}: no position on {when} for line {line.line_no}",
                        code="EVENT_WITHOUT_POSITION",
                    )
                event = CorporateEventEntity(
                    asset_id=asset_id,
                    kind=line.event_kind,  # type: ignore[arg-type]
                    ex_date=when,
                    factor=factor.quantize(Decimal("0.00000001")),
                    notes=note.format(line.line_no),
                    import_hash=line.hash,
                )
                events.append(event)
                new_events.append(event)
            try:
                validate_ledger(asset, ledger, events=events)
            except InsufficientQuantityError as err:
                raise LedgerConflictError(f"{ticker}: {err}") from err

            for trade in trades:
                uow.transactions.add(trade)
            for event in new_events:
                uow.corporate_events.add(event)
            for line in items:
                if line.action != LineAction.EARNING:
                    continue
                uow.earnings.add(
                    EarningEntity(
                        asset_id=asset_id,
                        kind=line.earning_kind,  # type: ignore[arg-type]
                        payment_date=line.when,  # type: ignore[arg-type]
                        gross_amount=line.gross_amount or Decimal("0"),
                        withholding_tax=line.withholding_tax or Decimal("0"),
                        notes=note.format(line.line_no),
                        account_id=account_id,
                        import_hash=line.hash,
                    )
                )
                created["EARNING"] += 1
            created["TRADE"] += len(trades)
            created["EVENT"] += len(new_events)
        return created
