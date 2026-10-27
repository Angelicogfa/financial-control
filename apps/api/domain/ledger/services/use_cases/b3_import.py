"""Use cases for the B3 statement import (preview and commit)."""

from __future__ import annotations

import csv
import io
from collections.abc import Callable, Iterable
from dataclasses import replace
from datetime import date, timedelta
from decimal import Decimal

from controllers.ledger.requests import ImportNewAssetRequest
from controllers.ledger.responses import (
    ImportCommitDTO,
    ImportFixedIncomeTermDTO,
    ImportLineDTO,
    ImportPreviewDTO,
    UnknownTickerDTO,
)
from controllers.shared.dependencies import UoWFactory
from core.errors import ValidationError
from domain.accounts.services.use_cases import resolve_account_id
from domain.assets.services.codes import generated_asset_code
from domain.ledger.services.b3_import import (
    FixedIncomeRef,
    ImportLine,
    ImportPreview,
    LineAction,
    LineStatus,
    classify,
    event_factor,
    guess_asset_type,
    mark_skipped,
    parse_statement,
    resolve_fixed_income,
)
from domain.ledger.services.ledger import LedgerConflictError, validate_ledger
from domain.portfolio.services.positions import InsufficientQuantityError, build_position
from domain.portfolio.services.snapshots import invalidate_snapshots_from
from domain.research.services.service import (
    register_researched_asset,
    suggested_arca_class,
)
from domain.shared.entities import (
    AssetEntity,
    CorporateEventEntity,
    EarningEntity,
    FixedIncomeTermEntity,
    TransactionEntity,
)
from domain.shared.enums import TransactionKind
from domain.shared.unit_of_work import AbstractUnitOfWork

MAX_STATEMENT_BYTES = 5 * 1024 * 1024


def read_statement_rows(filename: str, content: bytes) -> list[dict[str, object]]:
    """Rows of the first sheet (.xlsx) or of a CSV export, keyed by the header row."""
    if len(content) > MAX_STATEMENT_BYTES:
        raise ValidationError("Statement larger than 5 MB", code="FILE_TOO_LARGE")
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        except Exception as exc:
            raise ValidationError("Unreadable spreadsheet", code="INVALID_FILE") from exc
        # B3 exports declare a stale <dimension ref="A1"/>; read-only mode trusts it and would
        # return only the first cell, so scan the actual rows instead.
        sheet.reset_dimensions()
        values = [list(r) for r in sheet.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else ",", delimiters=";,\t")
        values = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ValidationError("Send the .xlsx (or .csv) exported by B3", code="INVALID_FILE")
    if not values:
        raise ValidationError("Empty statement", code="INVALID_FILE")
    headers = [str(h or "").strip() for h in values[0]]
    return [dict(zip(headers, row, strict=False)) for row in values[1:]]


def _parse(filename: str, content: bytes) -> ImportPreview:
    rows = read_statement_rows(filename, content)
    try:
        return parse_statement(rows)
    except ValueError as exc:
        raise ValidationError(str(exc), code="INVALID_STATEMENT") from exc


def _line_dto(line: ImportLine) -> ImportLineDTO:
    return ImportLineDTO(
        line_no=line.line_no,
        action=line.action.value,
        status=line.status.value,
        entry_date=line.when,
        ticker=line.ticker,
        description=line.description,
        quantity=line.quantity,
        unit_price=line.unit_price,
        amount=line.amount,
        trade_kind=line.trade_kind.value if line.trade_kind else None,
        earning_kind=line.earning_kind.value if line.earning_kind else None,
        event_kind=line.event_kind.value if line.event_kind else None,
        note=line.note,
    )


def _held_before(assets: dict[str, AssetEntity]) -> Callable[[str, date], Decimal]:
    def held(ticker: str, when: date) -> Decimal:
        asset = assets.get(ticker)
        if asset is None:
            return Decimal("0")
        try:
            return build_position(
                asset.transactions,
                up_to_date=when - timedelta(days=1),
                events=asset.corporate_events,
            ).quantity
        except InsufficientQuantityError:
            return Decimal("0")

    return held


def _fixed_income_refs(preview: ImportPreview) -> dict[str, tuple[FixedIncomeRef, date | None]]:
    """Group key -> (security, first date in the statement) of the lines with no asset."""
    refs: dict[str, tuple[FixedIncomeRef, date | None]] = {}
    for line in preview.lines:
        ref = line.fixed_income
        if ref is None or line.ticker != ref.key:
            continue
        _, first = refs.get(ref.key, (ref, line.when))
        refs[ref.key] = (ref, min((d for d in (first, line.when) if d), default=None))
    return refs


def _suggested_term(
    uow: AbstractUnitOfWork, ref: FixedIncomeRef, first: date | None
) -> ImportFixedIncomeTermDTO:
    maturity = None
    if ref.maturity_year is not None:
        maturity = next(
            (
                bond.maturity_date
                for bond in uow.treasury_catalog.list_all()
                if bond.instrument_kind == ref.instrument_kind
                and bond.maturity_date.year == ref.maturity_year
                and ("juros semestrais" in bond.bond_type.lower()) == ref.semiannual
            ),
            None,
        )
    return ImportFixedIncomeTermDTO(
        instrument_kind=ref.instrument_kind,
        indexer=ref.indexer,
        issue_date=first,
        maturity_date=maturity,
        issuer=ref.issuer,
    )


def _register_fixed_income(
    uow: AbstractUnitOfWork,
    ref: FixedIncomeRef,
    first: date | None,
    wanted: ImportNewAssetRequest,
) -> AssetEntity:
    """Register the Treasury bond or bank security with the terms chosen in the preview."""
    term = wanted.fixed_income_term
    if term is None or term.contracted_rate is None:
        raise ValidationError(
            f"{ref.name}: inform the contracted rate to register it",
            code="FIXED_INCOME_TERM_REQUIRED",
        )
    issue_date = term.issue_date or first
    if issue_date is None:
        raise ValidationError(
            f"{ref.name}: inform the issue date", code="FIXED_INCOME_TERM_REQUIRED"
        )
    entity_term = FixedIncomeTermEntity(
        instrument_kind=term.instrument_kind,
        indexer=term.indexer,
        contracted_rate=term.contracted_rate,
        issue_date=issue_date,
        maturity_date=term.maturity_date,
        tax_exempt=term.tax_exempt,
        liquidity=term.liquidity,
        issuer=term.issuer or ref.issuer,
    )
    name = wanted.name or ref.name
    ticker = generated_asset_code(
        ref.asset_type,
        name,
        entity_term,
        is_taken=lambda code: uow.assets.get_by_ticker(code) is not None,
    )
    return uow.assets.add(
        AssetEntity(
            ticker=ticker,
            name=name,
            asset_type=ref.asset_type,
            arca_class=wanted.arca_class or suggested_arca_class(ticker, ref.asset_type),
            fixed_income_term=entity_term,
        )
    )


UNBACKED_REDEMPTION = "Resgate maior que o saldo: a aplicação é anterior a este extrato"
UNBACKED_SELL = (
    "Venda maior que o saldo: a compra não está neste extrato nem nos lançamentos "
    "(ignore a linha ou lance a compra antes)"
)


def _reject_unbacked_sells(preview: ImportPreview, assets: dict[str, AssetEntity]) -> ImportPreview:
    """Mark INVALID the sells and redemptions larger than what the ledger holds.

    A statement covers a period: selling what was bought before it (or under another code,
    as after a conversion) has no position to come from. Rejecting the line keeps the rest
    of the statement importable, and the preview shows it before anything is written.
    """
    pending: dict[str, list[ImportLine]] = {}
    for line in preview.lines:
        if (
            line.action == LineAction.TRADE
            and line.status in (LineStatus.NEW, LineStatus.UNKNOWN_ASSET)
            and line.trade_kind is not None
            and line.when is not None
        ):
            pending.setdefault(line.ticker or "", []).append(line)
    rejected: dict[int, str] = {}
    for ticker, items in pending.items():
        asset = assets.get(ticker)
        if asset is None and any(i.fixed_income for i in items):
            continue  # a security is valued by its terms, chosen only at commit
        accepted = list(asset.transactions) if asset else []
        for line in sorted(
            items, key=lambda i: (i.when, i.trade_kind != TransactionKind.BUY, i.line_no)
        ):
            trade = TransactionEntity(
                asset_id=(asset.id if asset else None) or 0,
                kind=line.trade_kind,  # type: ignore[arg-type]
                trade_date=line.when,  # type: ignore[arg-type]
                quantity=line.quantity or Decimal("0"),
                unit_price=line.unit_price or Decimal("0"),
            )
            try:
                if asset is None:
                    build_position([*accepted, trade])
                else:
                    validate_ledger(asset, [*accepted, trade])
            except InsufficientQuantityError:
                rejected[line.line_no] = UNBACKED_REDEMPTION if line.fixed_income else UNBACKED_SELL
                continue
            accepted.append(trade)
    if not rejected:
        return preview
    return ImportPreview(
        preview.kind,
        [
            replace(line, status=LineStatus.INVALID, note=rejected[line.line_no])
            if line.line_no in rejected
            else line
            for line in preview.lines
        ],
    )


def _classified(
    uow: AbstractUnitOfWork,
    parsed: ImportPreview,
    assets: dict[str, AssetEntity],
    *,
    chosen: dict[str, str] | None = None,
    with_events: bool = False,
    skipped: Iterable[int] = (),
) -> ImportPreview:
    resolved = resolve_fixed_income(parsed, assets.values(), chosen=chosen)
    preview = classify(
        resolved,
        known_tickers=set(assets),
        existing_hashes=uow.ledger_imports.existing_hashes([line.hash for line in parsed.lines]),
        held_before=_held_before(assets) if with_events else None,
    )
    return _reject_unbacked_sells(mark_skipped(preview, skipped), assets)


class ImportHandler:
    """Preview shows what would happen; commit applies only NEW lines, all or nothing.

    `skip_lines` (statement line numbers) leaves lines out of both, so a sell whose purchase
    is missing, or any line the user does not want, never blocks the rest of the statement.
    """

    def __init__(self, uow_factory: UoWFactory) -> None:
        self.uow_factory = uow_factory

    def preview(
        self, filename: str, content: bytes, *, skip_lines: Iterable[int] = ()
    ) -> ImportPreviewDTO:
        parsed = _parse(filename, content)
        with self.uow_factory() as uow:
            assets = {a.ticker: a for a in uow.assets.list_all(active_only=False)}
            preview = _classified(uow, parsed, assets, with_events=True, skipped=skip_lines)
            refs = _fixed_income_refs(preview)
            unknown = []
            for ticker in preview.unknown_tickers:
                if ticker in refs:
                    ref, first = refs[ticker]
                    unknown.append(
                        UnknownTickerDTO(
                            ticker=ticker,
                            name=ref.name,
                            asset_type=ref.asset_type,
                            arca_class=suggested_arca_class(ticker, ref.asset_type),
                            fixed_income_term=_suggested_term(uow, ref, first),
                        )
                    )
                    continue
                listed = uow.catalog.get(ticker)
                asset_type = (listed.asset_type if listed else None) or guess_asset_type(ticker)
                unknown.append(
                    UnknownTickerDTO(
                        ticker=ticker,
                        name=listed.name if listed else None,
                        asset_type=asset_type,
                        arca_class=suggested_arca_class(ticker, asset_type) if asset_type else None,
                    )
                )
        return ImportPreviewDTO(
            kind=preview.kind.value,
            counts=preview.counts,
            unknown_tickers=unknown,
            lines=[_line_dto(line) for line in preview.lines],
        )

    def commit(
        self,
        filename: str,
        content: bytes,
        *,
        account_id: int | None,
        create_tickers: list[str],
        new_assets: Iterable[ImportNewAssetRequest] = (),
        skip_lines: Iterable[int] = (),
    ) -> ImportCommitDTO:
        """Register the chosen unknown tickers through the research flow, then apply the
        new lines. `create_tickers` takes the suggestions; `new_assets` carries choices.
        A ticker left with no line to import (all skipped or invalid) is not registered."""
        parsed = _parse(filename, content)
        skipped = list(skip_lines)
        chosen = {
            t.strip().upper(): ImportNewAssetRequest(ticker=t) for t in create_tickers if t.strip()
        }
        chosen |= {a.ticker.strip().upper(): a for a in new_assets}
        registered: list[str] = []
        fixed_income: dict[str, str] = {}
        with self.uow_factory() as uow:
            account = resolve_account_id(uow, account_id)
            existing = uow.assets.list_all(active_only=False)
            refs = _fixed_income_refs(resolve_fixed_income(parsed, existing))
            keys = {line.fixed_income.key for line in parsed.lines if line.fixed_income}
            needed = set(
                _classified(
                    uow, parsed, {a.ticker: a for a in existing}, skipped=skipped
                ).unknown_tickers
            )
            for ticker, wanted in sorted(chosen.items()):
                if ticker not in needed:
                    continue  # already registered, or nothing of it left to import
                if ticker in refs:
                    asset = _register_fixed_income(uow, *refs[ticker], wanted)
                    fixed_income[ticker] = asset.ticker
                    registered.append(asset.ticker)
                    continue
                if ticker in keys or uow.assets.get_by_ticker(ticker) is not None:
                    continue  # already registered (a security is matched, not looked up)
                listed = uow.catalog.get(ticker)
                asset_type = (
                    wanted.asset_type
                    or (listed.asset_type if listed else None)
                    or guess_asset_type(ticker)
                )
                if asset_type is None:
                    raise ValidationError(
                        f"Cannot tell the type of {ticker}; choose it in the preview",
                        code="UNKNOWN_INSTRUMENT",
                    )
                register_researched_asset(
                    uow, ticker, asset_type=asset_type, arca_class=wanted.arca_class
                )
                registered.append(ticker)

            assets = {a.ticker: a for a in uow.assets.list_all(active_only=False)}
            preview = _classified(uow, parsed, assets, chosen=fixed_income, skipped=skipped)
            created = self._apply(uow, preview.new_lines, assets, account)
            invalidate_snapshots_from(uow, *(line.when for line in preview.new_lines))
            uow.commit()
        return ImportCommitDTO(
            kind=preview.kind.value,
            created=created,
            skipped={k: v for k, v in preview.counts.items() if k != LineStatus.NEW.value},
            created_assets=registered,
        )

    @staticmethod
    def _apply(
        uow: AbstractUnitOfWork,
        lines: list[ImportLine],
        assets: dict[str, AssetEntity],
        account_id: int | None,
    ) -> dict[str, int]:
        created = {"TRADE": 0, "EARNING": 0, "EVENT": 0}
        by_ticker: dict[str, list[ImportLine]] = {}
        for line in lines:
            by_ticker.setdefault(line.ticker or "", []).append(line)

        for ticker, items in sorted(by_ticker.items()):
            asset = assets[ticker]
            asset_id = asset.id or 0
            note = "Importado da B3 (linha {})"
            trades = [
                TransactionEntity(
                    asset_id=asset_id,
                    kind=line.trade_kind,  # type: ignore[arg-type]
                    trade_date=line.when,  # type: ignore[arg-type]
                    quantity=line.quantity or Decimal("0"),
                    unit_price=line.unit_price or Decimal("0"),
                    notes=note.format(line.line_no),
                    account_id=account_id,
                    import_hash=line.hash,
                )
                for line in items
                if line.action == LineAction.TRADE
            ]
            ledger = [*asset.transactions, *trades]
            events = list(asset.corporate_events)
            new_events = []
            for line in sorted(
                (i for i in items if i.action == LineAction.EVENT),
                key=lambda i: (i.when, i.line_no),
            ):
                when = line.when or date.min
                try:
                    held = build_position(
                        ledger, up_to_date=when - timedelta(days=1), events=events
                    ).quantity
                except InsufficientQuantityError:
                    held = Decimal("0")
                factor = event_factor(line.event_kind, line.quantity or Decimal("0"), held)  # type: ignore[arg-type]
                if factor is None:
                    raise ValidationError(
                        f"{ticker}: no position on {when} for line {line.line_no}",
                        code="EVENT_WITHOUT_POSITION",
                    )
                event = CorporateEventEntity(
                    asset_id=asset_id,
                    kind=line.event_kind,  # type: ignore[arg-type]
                    ex_date=when,
                    factor=factor.quantize(Decimal("0.00000001")),
                    notes=note.format(line.line_no),
                    import_hash=line.hash,
                )
                events.append(event)
                new_events.append(event)
            try:
                validate_ledger(asset, ledger, events=events)
            except InsufficientQuantityError as err:
                raise LedgerConflictError(f"{ticker}: {err}") from err

            for trade in trades:
                uow.transactions.add(trade)
            for event in new_events:
                uow.corporate_events.add(event)
            for line in items:
                if line.action != LineAction.EARNING:
                    continue
                uow.earnings.add(
                    EarningEntity(
                        asset_id=asset_id,
                        kind=line.earning_kind,  # type: ignore[arg-type]
                        payment_date=line.when,  # type: ignore[arg-type]
                        gross_amount=line.gross_amount or Decimal("0"),
                        withholding_tax=line.withholding_tax or Decimal("0"),
                        notes=note.format(line.line_no),
                        account_id=account_id,
                        import_hash=line.hash,
                    )
                )
                created["EARNING"] += 1
            created["TRADE"] += len(trades)
            created["EVENT"] += len(new_events)
        return created
